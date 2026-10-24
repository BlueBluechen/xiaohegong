# -*- coding: utf-8 -*-
import os
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
HERE=os.path.dirname(os.path.abspath(__file__)); A=lambda f: os.path.join(HERE,"assets",f)
MM=1/25.4
SERIF="Songti SC"; SANS="PingFang SC"
D_TXT="F3F0E7"; D_GOLD="D9B96A"; D_MUT="7FA6A4"; CREAM="F5F1E7"; SEAL="C85A45"
def _f(r,s,fo,c,b=False,sp=None):
    r.font.size=Pt(s); r.font.bold=b; r.font.color.rgb=RGBColor.from_string(c); r.font.name=fo
    p=r._r.get_or_add_rPr()
    for t in ("ea","cs"):
        e=p.find(qn("a:"+t))
        if e is None: e=etree.SubElement(p,qn("a:"+t))
        e.set("typeface",fo)
    if sp is not None: p.set("spc",str(int(sp*100)))
AL={"l":PP_ALIGN.LEFT,"c":PP_ALIGN.CENTER,"r":PP_ALIGN.RIGHT}
AN={"t":MSO_ANCHOR.TOP,"m":MSO_ANCHOR.MIDDLE}
def txt(sl,x,y,w,h,ps,a="t"):
    if ps and "anchor" in ps[0]: a=ps[0]["anchor"]
    tb=sl.shapes.add_textbox(Inches(x*MM),Inches(y*MM),Inches(w*MM),Inches(h*MM))
    tf=tb.text_frame; tf.word_wrap=True; tf.auto_size=MSO_AUTO_SIZE.NONE
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=AN[a]
    for i,p in enumerate(ps):
        pa=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        pa.alignment=AL[p.get("al","l")]
        if p.get("lh"): pa.line_spacing=p["lh"]
        for t,d in p["seg"]:
            r=pa.add_run(); r.text=t
            _f(r,d.get("s",10),d.get("f",SANS),d.get("c",D_TXT),d.get("b",False),d.get("spc"))
def rect(sl,x,y,w,h,fill=None,line=None,rad=None,lw=0.75):
    sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rad else MSO_SHAPE.RECTANGLE,
        Inches(x*MM),Inches(y*MM),Inches(w*MM),Inches(h*MM))
    if rad:
        try: sh.adjustments[0]=rad
        except: pass
    if fill: sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor.from_string(fill)
    else: sh.fill.background()
    if line: sh.line.color.rgb=RGBColor.from_string(line); sh.line.width=Pt(lw)
    else: sh.line.fill.background()
    st=sh._element.find(qn("p:style"))
    if st is not None: sh._element.remove(st)
    sh.shadow.inherit=False
    return sh
def hline(sl,x,y,w,c,h=0.4): rect(sl,x,y,w,h,fill=c)

W,H=297,420
prs=Presentation(); prs.slide_width=Inches(W*MM); prs.slide_height=Inches(H*MM)
s=prs.slides.add_slide(prs.slide_layouts[6])
s.shapes.add_picture(A("poster_concept.png"),Inches(0),Inches(0),width=Inches(W*MM),height=Inches(H*MM))

# 顶部
txt(s,22,26,160,6,[{"seg":[("首站",{"s":9,"c":D_GOLD,"spc":3.2}),("　北京清河　│　2026",{"s":9,"c":D_MUT,"spc":3.2})]}])
hline(s,22,35,26,"2F6469",0.5)

# 主标题
txt(s,20,42,272,58,[{"seg":[("小小河工",{"s":150,"f":SERIF,"c":D_TXT,"spc":1.5})]}])
txt(s,22,113,240,10,[{"seg":[("AI+ 公民河流教育计划",{"s":20,"c":D_GOLD,"spc":4.2})]}])

# 河下方：主张
txt(s,22,288,255,34,[
  {"lh":1.42,"seg":[("一条河不需要更多赞美。",{"s":30,"f":SERIF,"c":D_TXT})]},
  {"lh":1.42,"seg":[("它需要更多记录。",{"s":30,"f":SERIF,"c":D_GOLD})]}])
txt(s,22,338,255,6,[{"seg":[("23.6 公里　·　十座水闸　·　四条支流　·　沿岸九十余种鸟类",{"s":9.5,"c":D_MUT,"spc":1.2})]}])

# 二维码
QX,QY,QS=W-22-32,336,32
s.shapes.add_picture(A("qr_demo.png"),Inches(QX*MM),Inches(QY*MM),width=Inches(QS*MM),height=Inches(QS*MM))
txt(s,QX-14,QY+QS+6,QS+14,5,[{"al":"r","seg":[("扫码开始巡河",{"s":8,"c":"9FC0C6","spc":1.0})]}])

# 页脚
hline(s,22,398,W-44,"1F5257",0.4)
txt(s,22,403,200,6,[{"seg":[("以河为目标　·　孩子是参与者　·　记录是产品",{"s":8,"c":"5E8C8E","spc":0.8})]}])
txt(s,W-22-90,403,90,6,[{"al":"r","seg":[("AI+ Civic River Education",{"s":7.5,"c":"4E7C82","spc":1.4})]}])

out=os.path.join(HERE,"小小河工-概念海报A3.pptx")
prs.save(out); print("saved",out)
